"""
Vista de Reporte Diario de Ventas
Muestra el resumen de ventas del día actual
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from typing import Dict, Any
from datetime import datetime
from views.base_view import BaseView
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os


class DailySalesReportView(BaseView):
    """Vista para mostrar el reporte diario de ventas"""
    
    def __init__(self, parent, controller, user_data: Dict[str, Any] = None, on_back=None):
        self.parent = parent
        self.controller = controller
        self.user_data = user_data or {}
        self.root = parent
        self.on_back = on_back
        
        # Datos del reporte actual
        self.current_report_data = None
        self.current_date = datetime.now().strftime('%Y-%m-%d')
        
        # Crear interfaz
        self.create_interface()
        
        # Cargar datos del día
        self.load_today_report()
    
    def create_interface(self):
        """Crear interfaz del reporte diario"""
        # Frame principal
        main_frame = tk.Frame(self.parent, bg='#f8f9fa')
        main_frame.pack(fill='both', expand=True)
        
        # Header
        self.create_header(main_frame)
        
        # Contenido
        self.create_content(main_frame)
    
    def create_header(self, parent):
        """Crear header del reporte"""
        header_frame = tk.Frame(parent, bg='#0ea5e9', height=120)
        header_frame.pack(fill='x')
        header_frame.pack_propagate(False)
        
        content_frame = tk.Frame(header_frame, bg='#0ea5e9')
        content_frame.pack(expand=True, fill='both', padx=30, pady=20)
        
        # Icono y título (izquierda)
        title_frame = tk.Frame(content_frame, bg='#0ea5e9')
        title_frame.pack(side='left')
        
        tk.Label(
            title_frame,
            text="📊",
            font=('Segoe UI', 36),
            bg='#0ea5e9',
            fg='white'
        ).pack(side='left', padx=(0, 15))
        
        tk.Label(
            title_frame,
            text="Reporte Diario de Ventas",
            font=('Segoe UI', 24, 'bold'),
            bg='#0ea5e9',
            fg='white'
        ).pack(side='left')
        
        # Botón de volver (derecha)
        if self.on_back:
            tk.Button(
                content_frame,
                text="⬅️ Volver",
                command=self.back_to_dashboard,
                bg='#0284c7',
                fg='white',
                font=('Segoe UI', 11, 'bold'),
                relief='flat',
                cursor='hand2',
                padx=20,
                pady=8
            ).pack(side='right', padx=(20, 0))
        
        # Fecha actual (derecha, antes del botón)
        today = datetime.now()
        date_text = today.strftime("%A, %d de %B de %Y")
        # Traducir días y meses
        days = {'Monday': 'Lunes', 'Tuesday': 'Martes', 'Wednesday': 'Miércoles',
                'Thursday': 'Jueves', 'Friday': 'Viernes', 'Saturday': 'Sábado', 'Sunday': 'Domingo'}
        months = {'January': 'Enero', 'February': 'Febrero', 'March': 'Marzo', 'April': 'Abril',
                  'May': 'Mayo', 'June': 'Junio', 'July': 'Julio', 'August': 'Agosto',
                  'September': 'Septiembre', 'October': 'Octubre', 'November': 'Noviembre', 'December': 'Diciembre'}
        
        for en, es in days.items():
            date_text = date_text.replace(en, es)
        for en, es in months.items():
            date_text = date_text.replace(en, es)
        
        self.date_label = tk.Label(
            content_frame,
            text=date_text,
            font=('Segoe UI', 14),
            bg='#0ea5e9',
            fg='white'
        )
        self.date_label.pack(side='right', padx=(0, 15))
    
    def create_content(self, parent):
        """Crear contenido del reporte"""
        content_frame = tk.Frame(parent, bg='#f8f9fa')
        content_frame.pack(fill='both', expand=True, padx=30, pady=30)
        
        # Indicador de loading (se mostrará al inicio)
        self.loading_frame = tk.Frame(content_frame, bg='#f8f9fa')
        self.loading_frame.pack(fill='both', expand=True)
        
        tk.Label(
            self.loading_frame,
            text="⏳",
            font=('Segoe UI', 48),
            bg='#f8f9fa',
            fg='#3498db'
        ).pack(pady=(100, 20))
        
        tk.Label(
            self.loading_frame,
            text="Cargando datos del día...",
            font=('Segoe UI', 16, 'bold'),
            bg='#f8f9fa',
            fg='#7f8c8d'
        ).pack()
        
        # Frame para el contenido real (se mostrará después de cargar)
        self.content_main_frame = tk.Frame(content_frame, bg='#f8f9fa')
        
        # Grid para tarjetas de resumen (solo 3 tarjetas)
        cards_frame = tk.Frame(self.content_main_frame, bg='#f8f9fa')
        cards_frame.pack(fill='x', pady=(0, 30))
        
        # Tarjeta 1: Total Ventas (centrada)
        self.create_summary_card(
            cards_frame,
            "💰 Total Ventas",
            "S/ 0.00",
            "#27ae60",
            0, 0
        )
        
        # Tarjeta 2: Número de Ventas (centrada)
        self.create_summary_card(
            cards_frame,
            "🛒 Número de Ventas",
            "0",
            "#3498db",
            0, 1
        )
        
        # Tarjeta 3: Productos Vendidos (centrada)
        self.create_summary_card(
            cards_frame,
            "📦 Productos Vendidos",
            "0",
            "#e67e22",
            0, 2
        )
        
        # Configurar grid - solo 3 columnas
        for i in range(3):
            cards_frame.columnconfigure(i, weight=1, uniform='card')
        
        # Tabla de ventas del día
        self.create_sales_table(self.content_main_frame)
        
        # Botones de acción
        buttons_frame = tk.Frame(self.content_main_frame, bg='#f8f9fa')
        buttons_frame.pack(fill='x', pady=(20, 0))
        
        tk.Button(
            buttons_frame,
            text="🔄 Actualizar",
            command=self.load_today_report,
            bg='#3498db',
            fg='white',
            font=('Segoe UI', 13, 'bold'),
            relief='flat',
            cursor='hand2',
            padx=25,
            pady=12
        ).pack(side='left', padx=(0, 15))
        
        tk.Button(
            buttons_frame,
            text="📊 Exportar a Excel",
            command=self.export_report,
            bg='#27ae60',
            fg='white',
            font=('Segoe UI', 13, 'bold'),
            relief='flat',
            cursor='hand2',
            padx=25,
            pady=12
        ).pack(side='left')
    
    def create_summary_card(self, parent, title, value, color, row, col):
        """Crear tarjeta de resumen"""
        card = tk.Frame(parent, bg='white', relief='solid', bd=1)
        card.grid(row=row, column=col, padx=10, pady=10, sticky='nsew')
        
        # Padding interno
        inner = tk.Frame(card, bg='white')
        inner.pack(fill='both', expand=True, padx=25, pady=25)
        
        # Barra de color
        tk.Frame(inner, bg=color, height=4).pack(fill='x', pady=(0, 15))
        
        # Título
        tk.Label(
            inner,
            text=title,
            font=('Segoe UI', 12, 'bold'),
            bg='white',
            fg='#7f8c8d'
        ).pack(anchor='w')
        
        # Valor
        value_label = tk.Label(
            inner,
            text=value,
            font=('Segoe UI', 28, 'bold'),
            bg='white',
            fg='#2c3e50'
        )
        value_label.pack(anchor='w', pady=(10, 0))
        
        # Guardar referencia
        card.value_label = value_label
        
        # Guardar referencia en la vista
        card_key = title.split()[1].lower() if len(title.split()) > 1 else title.lower()
        setattr(self, f'card_{card_key}', card)
        print(f"   📝 Tarjeta creada: card_{card_key} = '{title}'")
    
    def create_sales_table(self, parent):
        """Crear tabla de ventas del día"""
        # Frame contenedor
        table_frame = tk.Frame(parent, bg='white', relief='solid', bd=1)
        table_frame.pack(fill='both', expand=True, padx=10, pady=(20, 10))
        
        # Título
        title_frame = tk.Frame(table_frame, bg='#3498db', height=40)
        title_frame.pack(fill='x')
        title_frame.pack_propagate(False)
        
        tk.Label(
            title_frame,
            text="📋 Ventas del Día",
            font=('Segoe UI', 14, 'bold'),
            bg='#3498db',
            fg='white'
        ).pack(side='left', padx=15, pady=8)
        
        # Frame para la tabla con scrollbars
        tree_frame = tk.Frame(table_frame, bg='white')
        tree_frame.pack(fill='both', expand=True, padx=2, pady=2)
        
        # Scrollbars
        vsb = ttk.Scrollbar(tree_frame, orient="vertical")
        vsb.pack(side='right', fill='y')
        
        hsb = ttk.Scrollbar(tree_frame, orient="horizontal")
        hsb.pack(side='bottom', fill='x')
        
        # Treeview
        columns = ('ID', 'Hora', 'Cajero', 'Total', 'Método Pago', 'Estado')
        self.sales_tree = ttk.Treeview(
            tree_frame,
            columns=columns,
            show='headings',
            yscrollcommand=vsb.set,
            xscrollcommand=hsb.set,
            height=10
        )
        
        vsb.config(command=self.sales_tree.yview)
        hsb.config(command=self.sales_tree.xview)
        
        # Configurar columnas
        self.sales_tree.heading('ID', text='ID')
        self.sales_tree.heading('Hora', text='Hora')
        self.sales_tree.heading('Cajero', text='Cajero')
        self.sales_tree.heading('Total', text='Total')
        self.sales_tree.heading('Método Pago', text='Método de Pago')
        self.sales_tree.heading('Estado', text='Estado')
        
        self.sales_tree.column('ID', width=60, anchor='center')
        self.sales_tree.column('Hora', width=100, anchor='center')
        self.sales_tree.column('Cajero', width=150, anchor='w')
        self.sales_tree.column('Total', width=120, anchor='e')
        self.sales_tree.column('Método Pago', width=150, anchor='center')
        self.sales_tree.column('Estado', width=120, anchor='center')
        
        # Estilo alternado de filas
        self.sales_tree.tag_configure('oddrow', background='#f8f9fa')
        self.sales_tree.tag_configure('evenrow', background='white')
        
        self.sales_tree.pack(fill='both', expand=True)
    
    def update_sales_table(self, sales):
        """Actualizar la tabla de ventas con los datos"""
        print(f"   📋 Actualizando tabla con {len(sales)} ventas")
        
        # Limpiar tabla
        for item in self.sales_tree.get_children():
            self.sales_tree.delete(item)
        
        # Traducción de métodos de pago
        payment_methods = {
            'cash': 'Efectivo',
            'card': 'Tarjeta',
            'transfer': 'Transferencia',
            'yape': 'Yape',
            'plin': 'Plin'
        }
        
        # Traducción de estados
        status_translation = {
            'completed': 'Completado',
            'cancelled': 'Cancelado',
            'pending': 'Pendiente'
        }
        
        # Insertar ventas
        for idx, sale in enumerate(sales):
            # Extraer hora de la fecha
            sale_datetime = sale.get('sale_date', '')
            if isinstance(sale_datetime, str):
                try:
                    time_str = sale_datetime.split(' ')[1] if ' ' in sale_datetime else sale_datetime
                except:
                    time_str = sale_datetime
            else:
                time_str = str(sale_datetime)
            
            # Formatear datos
            sale_id = sale.get('id', '')
            cashier = sale.get('cashier_name', 'N/A')
            total = f"S/ {sale.get('total_amount', 0):,.2f}"
            payment = payment_methods.get(sale.get('payment_method', ''), sale.get('payment_method', 'N/A'))
            status = status_translation.get(sale.get('status', ''), sale.get('status', 'N/A'))
            
            # Tag para fila alternada
            tag = 'evenrow' if idx % 2 == 0 else 'oddrow'
            
            # Insertar en tabla
            self.sales_tree.insert(
                '',
                'end',
                values=(sale_id, time_str, cashier, total, payment, status),
                tags=(tag,)
            )
        
        print(f"   ✅ Tabla actualizada con {len(sales)} ventas")
    
    def load_today_report(self):
        """Cargar reporte del día actual con indicador de loading"""
        try:
            print("📊 DEBUG: Iniciando carga de reporte diario")
            
            # Mostrar loading
            if hasattr(self, 'content_main_frame'):
                self.content_main_frame.pack_forget()
                print("   - content_main_frame ocultado")
            if hasattr(self, 'loading_frame'):
                self.loading_frame.pack(fill='both', expand=True)
                print("   - loading_frame mostrado")
            
            # Actualizar la interfaz para mostrar el loading
            self.root.update_idletasks()
            
            # Obtener fecha de hoy
            today = datetime.now().strftime('%Y-%m-%d')
            self.current_date = today
            print(f"   - Fecha: {today}")
            
            # Obtener resumen del día
            result = self.controller.get_daily_summary(today)
            print(f"   - Resultado success: {result.get('success')}")
            print(f"   - Datos: {result.get('data')}")
            
            if result.get('success'):
                data = result.get('data', {})
                self.current_report_data = data
                
                # Los datos vienen en data['sales_summary']
                summary = data.get('sales_summary', {})
                print(f"   - Summary completo: {summary}")
                
                # Actualizar solo las 3 tarjetas necesarias
                if hasattr(self, 'card_total'):
                    total = summary.get('total_amount', 0)
                    self.card_total.value_label.config(text=f"S/ {total:,.2f}")
                    print(f"   ✅ Total Ventas: S/ {total:,.2f}")
                else:
                    print(f"   ❌ card_total NO existe")
                
                if hasattr(self, 'card_número'):
                    count = summary.get('total_sales', 0)
                    self.card_número.value_label.config(text=str(count))
                    print(f"   ✅ Número de Ventas: {count}")
                else:
                    print(f"   ❌ card_número NO existe")
                
                if hasattr(self, 'card_productos'):
                    # Obtener total de productos vendidos
                    products = summary.get('total_products', 0)
                    self.card_productos.value_label.config(text=str(products))
                    print(f"   ✅ Productos Vendidos: {products}")
                else:
                    print(f"   ❌ card_productos NO existe")
                
                # Actualizar tabla de ventas
                self.update_sales_table(data.get('sales', []))
                
                # Ocultar loading y mostrar contenido
                if hasattr(self, 'loading_frame'):
                    self.loading_frame.pack_forget()
                    print("   - loading_frame ocultado")
                if hasattr(self, 'content_main_frame'):
                    self.content_main_frame.pack(fill='both', expand=True)
                    print("   - content_main_frame mostrado")
                    
            else:
                # Ocultar loading
                if hasattr(self, 'loading_frame'):
                    self.loading_frame.pack_forget()
                if hasattr(self, 'content_main_frame'):
                    self.content_main_frame.pack(fill='both', expand=True)
                    
                self.current_report_data = None
                print("   ⚠️ No hay datos de ventas")
                messagebox.showwarning(
                    "Sin Datos",
                    "No hay ventas registradas para el día de hoy."
                )
                
        except Exception as e:
            # Ocultar loading en caso de error
            if hasattr(self, 'loading_frame'):
                self.loading_frame.pack_forget()
            if hasattr(self, 'content_main_frame'):
                self.content_main_frame.pack(fill='both', expand=True)
                
            messagebox.showerror(
                "Error",
                f"Error al cargar el reporte: {str(e)}"
            )
            print(f"❌ ERROR en load_today_report: {e}")
            import traceback
            traceback.print_exc()
    
    def view_detail(self):
        """Ver detalle de ventas del día"""
        if not self.current_report_data:
            messagebox.showwarning(
                "Sin Datos",
                "No hay datos para mostrar."
            )
            return
        
        try:
            # Obtener ventas detalladas del día
            result = self.controller.get_sales_report(
                self.current_date,
                self.current_date
            )
            
            if not result.get('success'):
                messagebox.showerror(
                    "Error",
                    result.get('message', 'No se pudieron obtener los detalles')
                )
                return
            
            sales_data = result.get('data', {})
            sales_list = sales_data.get('sales', [])
            
            if not sales_list:
                messagebox.showinfo(
                    "Sin Ventas",
                    "No hay ventas registradas para este día."
                )
                return
            
            # Crear ventana de detalle
            detail_window = tk.Toplevel(self.root)
            detail_window.title("Detalle de Ventas del Día")
            detail_window.geometry("1000x600")
            detail_window.transient(self.root)
            detail_window.grab_set()
            
            # Header
            header = tk.Frame(detail_window, bg='#3498db', height=80)
            header.pack(fill='x')
            header.pack_propagate(False)
            
            tk.Label(
                header,
                text=f"📋 Detalle de Ventas - {self._format_date_spanish(self.current_date)}",
                font=('Segoe UI', 18, 'bold'),
                bg='#3498db',
                fg='white'
            ).pack(pady=20)
            
            # Frame principal
            main_frame = tk.Frame(detail_window, bg='white')
            main_frame.pack(fill='both', expand=True, padx=20, pady=20)
            
            # Resumen rápido
            summary_frame = tk.Frame(main_frame, bg='#ecf0f1', relief='solid', bd=1)
            summary_frame.pack(fill='x', pady=(0, 15))
            
            summary_inner = tk.Frame(summary_frame, bg='#ecf0f1')
            summary_inner.pack(fill='x', padx=20, pady=15)
            
            summary_data = sales_data.get('summary', {})
            
            tk.Label(
                summary_inner,
                text=f"Total Ventas: {summary_data.get('total_sales', 0)}",
                font=('Segoe UI', 12, 'bold'),
                bg='#ecf0f1',
                fg='#2c3e50'
            ).pack(side='left', padx=15)
            
            tk.Label(
                summary_inner,
                text=f"Monto Total: S/ {summary_data.get('total_amount', 0):,.2f}",
                font=('Segoe UI', 12, 'bold'),
                bg='#ecf0f1',
                fg='#27ae60'
            ).pack(side='left', padx=15)
            
            tk.Label(
                summary_inner,
                text=f"Ticket Promedio: S/ {summary_data.get('average_ticket', 0):,.2f}",
                font=('Segoe UI', 12, 'bold'),
                bg='#ecf0f1',
                fg='#9b59b6'
            ).pack(side='left', padx=15)
            
            # Tabla de ventas
            table_frame = tk.Frame(main_frame, bg='white')
            table_frame.pack(fill='both', expand=True)
            
            # Scrollbars
            v_scroll = ttk.Scrollbar(table_frame, orient='vertical')
            h_scroll = ttk.Scrollbar(table_frame, orient='horizontal')
            
            # Treeview
            columns = ('ID', 'Hora', 'Cajero', 'Total', 'Pago', 'Estado')
            tree = ttk.Treeview(
                table_frame,
                columns=columns,
                show='headings',
                yscrollcommand=v_scroll.set,
                xscrollcommand=h_scroll.set,
                height=15
            )
            
            v_scroll.config(command=tree.yview)
            h_scroll.config(command=tree.xview)
            
            # Configurar columnas
            tree.heading('ID', text='ID')
            tree.heading('Hora', text='Hora')
            tree.heading('Cajero', text='Cajero')
            tree.heading('Total', text='Total')
            tree.heading('Pago', text='Método Pago')
            tree.heading('Estado', text='Estado')
            
            tree.column('ID', width=80, anchor='center')
            tree.column('Hora', width=150, anchor='center')
            tree.column('Cajero', width=200, anchor='w')
            tree.column('Total', width=120, anchor='e')
            tree.column('Pago', width=150, anchor='center')
            tree.column('Estado', width=100, anchor='center')
            
            # Insertar datos
            for sale in sales_list:
                sale_id = sale.get('id', 'N/A')
                sale_time = sale.get('created_at', '')
                if isinstance(sale_time, str):
                    try:
                        dt = datetime.strptime(sale_time, '%Y-%m-%d %H:%M:%S')
                        sale_time = dt.strftime('%H:%M:%S')
                    except:
                        pass
                elif hasattr(sale_time, 'strftime'):
                    sale_time = sale_time.strftime('%H:%M:%S')
                
                cashier = sale.get('user_name', 'N/A')
                total = f"S/ {sale.get('total', 0):,.2f}"
                payment = sale.get('payment_method', 'N/A')
                
                payment_names = {
                    'cash': 'Efectivo',
                    'card': 'Tarjeta',
                    'yape': 'Yape',
                    'plin': 'Plin',
                    'transfer': 'Transferencia'
                }
                payment = payment_names.get(payment, payment)
                
                status = sale.get('status', 'completed')
                status_names = {
                    'completed': 'Completado',
                    'cancelled': 'Cancelado',
                    'pending': 'Pendiente'
                }
                status = status_names.get(status, status)
                
                tree.insert('', 'end', values=(
                    sale_id, sale_time, cashier, total, payment, status
                ))
            
            # Empaquetar tabla
            tree.pack(side='left', fill='both', expand=True)
            v_scroll.pack(side='right', fill='y')
            h_scroll.pack(side='bottom', fill='x')
            
            # Botón cerrar
            btn_frame = tk.Frame(detail_window, bg='white')
            btn_frame.pack(fill='x', padx=20, pady=(0, 20))
            
            tk.Button(
                btn_frame,
                text="Cerrar",
                command=detail_window.destroy,
                bg='#95a5a6',
                fg='white',
                font=('Segoe UI', 12, 'bold'),
                relief='flat',
                cursor='hand2',
                padx=30,
                pady=10
            ).pack(side='right')
            
        except Exception as e:
            messagebox.showerror(
                "Error",
                f"Error al mostrar detalle: {str(e)}"
            )
            print(f"Error en view_detail: {e}")
            import traceback
            traceback.print_exc()
    
    def export_report(self):
        """Exportar reporte a Excel"""
        if not self.current_report_data:
            messagebox.showwarning(
                "Sin Datos",
                "No hay datos para exportar."
            )
            return
        
        try:
            # Obtener ventas detalladas
            result = self.controller.get_sales_report(
                self.current_date,
                self.current_date
            )
            
            if not result.get('success'):
                messagebox.showerror(
                    "Error",
                    result.get('message', 'No se pudieron obtener los datos')
                )
                return
            
            sales_data = result.get('data', {})
            sales_list = sales_data.get('sales', [])
            summary = sales_data.get('summary', {})
            
            # Pedir ubicación de guardado
            fecha_formateada = datetime.now().strftime('%Y%m%d')
            default_filename = f"Reporte_Diario_{fecha_formateada}.xlsx"
            
            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                initialfile=default_filename,
                title="Guardar Reporte"
            )
            
            if not file_path:
                return  # Usuario canceló
            
            # Crear workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reporte Diario"
            
            # Estilos
            header_fill = PatternFill(start_color="3498db", end_color="3498db", fill_type="solid")
            header_font = Font(name='Calibri', size=12, bold=True, color="FFFFFF")
            title_font = Font(name='Calibri', size=16, bold=True, color="2c3e50")
            subtitle_font = Font(name='Calibri', size=11, bold=True, color="7f8c8d")
            normal_font = Font(name='Calibri', size=11)
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Título
            ws.merge_cells('A1:F1')
            ws['A1'] = 'REPORTE DIARIO DE VENTAS'
            ws['A1'].font = title_font
            ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
            
            # Fecha
            ws.merge_cells('A2:F2')
            ws['A2'] = f'Fecha: {self._format_date_spanish(self.current_date)}'
            ws['A2'].font = subtitle_font
            ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
            
            # Espacio
            ws.append([])
            
            # Resumen
            ws.merge_cells('A4:F4')
            ws['A4'] = 'RESUMEN DEL DÍA'
            ws['A4'].font = Font(name='Calibri', size=13, bold=True, color="2c3e50")
            ws['A4'].alignment = Alignment(horizontal='center', vertical='center')
            ws['A4'].fill = PatternFill(start_color="ecf0f1", end_color="ecf0f1", fill_type="solid")
            
            # Datos de resumen
            summary_row = 5
            ws[f'A{summary_row}'] = 'Total Ventas:'
            ws[f'B{summary_row}'] = summary.get('total_sales', 0)
            ws[f'D{summary_row}'] = 'Monto Total:'
            ws[f'E{summary_row}'] = f"S/ {summary.get('total_amount', 0):,.2f}"
            
            summary_row += 1
            ws[f'A{summary_row}'] = 'Ticket Promedio:'
            ws[f'B{summary_row}'] = f"S/ {summary.get('average_ticket', 0):,.2f}"
            ws[f'D{summary_row}'] = 'Total Productos:'
            ws[f'E{summary_row}'] = self.current_report_data.get('total_products', 0)
            
            # Aplicar estilos a resumen
            for row in range(5, 7):
                for col in ['A', 'B', 'D', 'E']:
                    cell = ws[f'{col}{row}']
                    cell.font = normal_font
                    cell.border = border
                    if col in ['A', 'D']:
                        cell.font = Font(name='Calibri', size=11, bold=True)
            
            # Espacio
            ws.append([])
            ws.append([])
            
            # Detalle de ventas
            ws.merge_cells('A9:F9')
            ws['A9'] = 'DETALLE DE VENTAS'
            ws['A9'].font = Font(name='Calibri', size=13, bold=True, color="2c3e50")
            ws['A9'].alignment = Alignment(horizontal='center', vertical='center')
            ws['A9'].fill = PatternFill(start_color="ecf0f1", end_color="ecf0f1", fill_type="solid")
            
            # Cabeceras de tabla
            headers = ['ID', 'Hora', 'Cajero', 'Total', 'Método Pago', 'Estado']
            header_row = 10
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=col)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            # Datos de ventas
            payment_names = {
                'cash': 'Efectivo',
                'card': 'Tarjeta',
                'yape': 'Yape',
                'plin': 'Plin',
                'transfer': 'Transferencia'
            }
            
            status_names = {
                'completed': 'Completado',
                'cancelled': 'Cancelado',
                'pending': 'Pendiente'
            }
            
            data_row = header_row + 1
            for sale in sales_list:
                sale_time = sale.get('created_at', '')
                if isinstance(sale_time, str):
                    try:
                        dt = datetime.strptime(sale_time, '%Y-%m-%d %H:%M:%S')
                        sale_time = dt.strftime('%H:%M:%S')
                    except:
                        pass
                elif hasattr(sale_time, 'strftime'):
                    sale_time = sale_time.strftime('%H:%M:%S')
                
                ws.cell(row=data_row, column=1, value=sale.get('id', 'N/A'))
                ws.cell(row=data_row, column=2, value=sale_time)
                ws.cell(row=data_row, column=3, value=sale.get('user_name', 'N/A'))
                ws.cell(row=data_row, column=4, value=sale.get('total', 0))
                ws.cell(row=data_row, column=5, value=payment_names.get(sale.get('payment_method', 'N/A'), sale.get('payment_method', 'N/A')))
                ws.cell(row=data_row, column=6, value=status_names.get(sale.get('status', 'completed'), sale.get('status', 'completed')))
                
                # Aplicar estilos
                for col in range(1, 7):
                    cell = ws.cell(row=data_row, column=col)
                    cell.font = normal_font
                    cell.border = border
                    cell.alignment = Alignment(horizontal='center' if col in [1, 2, 5, 6] else 'left')
                    
                    # Formato de moneda para total
                    if col == 4:
                        cell.number_format = '"S/ "#,##0.00'
                        cell.alignment = Alignment(horizontal='right')
                
                data_row += 1
            
            # Ajustar anchos de columna
            ws.column_dimensions['A'].width = 10
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 20
            ws.column_dimensions['F'].width = 15
            
            # Guardar
            wb.save(file_path)
            
            messagebox.showinfo(
                "Éxito",
                f"Reporte exportado exitosamente a:\n{file_path}"
            )
            
            # Preguntar si desea abrir
            if messagebox.askyesno("Abrir Archivo", "¿Desea abrir el archivo ahora?"):
                os.startfile(file_path)
                
        except Exception as e:
            messagebox.showerror(
                "Error",
                f"Error al exportar reporte: {str(e)}"
            )
            print(f"Error en export_report: {e}")
            import traceback
            traceback.print_exc()
    
    def _format_date_spanish(self, date_str):
        """Formatear fecha en español"""
        try:
            dt = datetime.strptime(date_str, '%Y-%m-%d')
            
            days = {
                'Monday': 'Lunes', 'Tuesday': 'Martes', 'Wednesday': 'Miércoles',
                'Thursday': 'Jueves', 'Friday': 'Viernes', 'Saturday': 'Sábado',
                'Sunday': 'Domingo'
            }
            
            months = {
                'January': 'Enero', 'February': 'Febrero', 'March': 'Marzo',
                'April': 'Abril', 'May': 'Mayo', 'June': 'Junio',
                'July': 'Julio', 'August': 'Agosto', 'September': 'Septiembre',
                'October': 'Octubre', 'November': 'Noviembre', 'December': 'Diciembre'
            }
            
            date_text = dt.strftime("%A, %d de %B de %Y")
            
            for en, es in days.items():
                date_text = date_text.replace(en, es)
            for en, es in months.items():
                date_text = date_text.replace(en, es)
            
            return date_text
        except:
            return date_str
    
    def back_to_dashboard(self):
        """Volver al dashboard"""
        if self.on_back:
            self.on_back()
        else:
            messagebox.showinfo(
                "Información",
                "No se ha configurado la navegación al dashboard."
            )
