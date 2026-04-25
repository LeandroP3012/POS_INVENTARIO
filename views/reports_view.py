"""
Vista de Reportes Avanzados
Autor: Sistema POS
Fecha: 2025-10-27
Actualizado: 2026-02-11 - Reportes avanzados
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, timedelta
from tkcalendar import DateEntry
import os
import threading
from collections import defaultdict

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure


class ReportsView:
    """Vista para módulo de reportes"""
    
    def __init__(self, parent, controller, user_data, on_back=None):
        self.parent = parent
        self.controller = controller
        self.user_data = user_data
        self.on_back = on_back
        
        # Frame principal
        self.main_frame = tk.Frame(parent, bg='#ecf0f1')
        self.main_frame.pack(fill='both', expand=True)
        
        self.setup_ui()
    
    def setup_ui(self):
        """Configurar interfaz"""
        self.create_header()
        
        content = tk.Frame(self.main_frame, bg='#ecf0f1')
        content.pack(fill='both', expand=True, padx=10, pady=10)
        
        # Panel izquierdo: Opciones
        left_panel = tk.Frame(content, bg='white', relief='solid', borderwidth=1, width=300)
        left_panel.pack(side='left', fill='y', padx=(0, 5))
        left_panel.pack_propagate(False)
        self.create_report_options(left_panel)
        
        # Panel derecho: Visualización con scroll
        right_outer = tk.Frame(content, bg='white', relief='solid', borderwidth=1)
        right_outer.pack(side='left', fill='both', expand=True, padx=(5, 0))
        
        self._right_canvas = tk.Canvas(right_outer, bg='white', highlightthickness=0)
        self._right_scrollbar = tk.Scrollbar(right_outer, orient='vertical', command=self._right_canvas.yview)
        self._right_scrollbar.pack(side='right', fill='y')
        self._right_canvas.pack(side='left', fill='both', expand=True)
        self._right_canvas.configure(yscrollcommand=self._right_scrollbar.set)
        
        self.right_panel = tk.Frame(self._right_canvas, bg='white')
        self._right_canvas_window = self._right_canvas.create_window(
            (0, 0), window=self.right_panel, anchor='nw'
        )
        
        # Ajustar scroll region y ancho del frame interno
        def _configure_scroll(event=None):
            self._right_canvas.configure(scrollregion=self._right_canvas.bbox('all'))
        
        def _configure_width(event):
            self._right_canvas.itemconfig(self._right_canvas_window, width=event.width)
        
        self.right_panel.bind('<Configure>', _configure_scroll)
        self._right_canvas.bind('<Configure>', _configure_width)
        
        # Scroll con rueda del mouse
        def _on_mousewheel(event):
            self._right_canvas.yview_scroll(int(-1 * (event.delta / 120)), 'units')
        
        self._right_canvas.bind_all('<MouseWheel>', _on_mousewheel)
        
        self.create_welcome_screen()
    
    def create_header(self):
        """Crear header"""
        header = tk.Frame(self.main_frame, bg='#8e44ad', height=80)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(
            header, text="📊 REPORTES Y ANÁLISIS",
            font=('Segoe UI', 20, 'bold'), bg='#8e44ad', fg='white'
        ).pack(side='left', padx=20, pady=10)
        
        user_info = tk.Frame(header, bg='#8e44ad')
        user_info.pack(side='right', padx=20, pady=10)
        
        tk.Label(
            user_info,
            text=f"👤 {self.user_data.get('full_name', 'Usuario')}",
            font=('Segoe UI', 10), bg='#8e44ad', fg='white'
        ).pack()
        
        tk.Label(
            user_info,
            text=f"📅 {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            font=('Segoe UI', 9), bg='#8e44ad', fg='#ecf0f1'
        ).pack()
        
        if self.on_back:
            tk.Button(
                header, text="⬅️ Volver", command=self.on_back,
                font=('Segoe UI', 10), bg='#9b59b6', fg='white',
                relief='flat', padx=15, pady=5, cursor='hand2'
            ).pack(side='right', padx=10)
    
    def create_report_options(self, parent):
        """Crear opciones de reportes"""
        tk.Label(
            parent, text="TIPOS DE REPORTES",
            font=('Segoe UI', 14, 'bold'), bg='white', fg='#2c3e50'
        ).pack(pady=15)
        
        reports = [
            {'name': 'Reporte de Ventas', 'icon': '📈', 'command': self.show_sales_report,
             'desc': 'Ventas por período'},
            {'name': 'Flujo de Productos', 'icon': '📦', 'command': self.show_products_report,
             'desc': 'Entradas y salidas'},
            {'name': 'Ventas Mensuales', 'icon': '📅', 'command': self.show_monthly_sales_report,
             'desc': 'Resumen mes a mes'},
            {'name': 'Inversión Inventario', 'icon': '💰', 'command': self.show_inventory_report,
             'desc': 'Capital invertido en tienda'},
            {'name': 'Producto Específico', 'icon': '🔍', 'command': self.show_specific_product_report,
             'desc': 'Rendimiento individual'},
        ]
        
        for report in reports:
            frame = tk.Frame(parent, bg='white')
            frame.pack(fill='x', padx=10, pady=5)
            
            btn = tk.Button(
                frame, text=f"{report['icon']} {report['name']}",
                command=report['command'],
                font=('Segoe UI', 11, 'bold'), bg='#8e44ad', fg='white',
                relief='flat', padx=15, pady=10, cursor='hand2', anchor='w'
            )
            btn.pack(fill='x')
            btn.bind('<Enter>', lambda e, b=btn: b.config(bg='#9b59b6'))
            btn.bind('<Leave>', lambda e, b=btn: b.config(bg='#8e44ad'))
            
            tk.Label(
                frame, text=report['desc'],
                font=('Segoe UI', 8), bg='white', fg='#7f8c8d'
            ).pack(anchor='w', padx=5)
    
    def create_welcome_screen(self):
        """Pantalla de bienvenida"""
        for widget in self.right_panel.winfo_children():
            widget.destroy()
        
        welcome = tk.Frame(self.right_panel, bg='white')
        welcome.pack(expand=True)
        
        tk.Label(welcome, text="📊", font=('Segoe UI', 72), bg='white').pack(pady=20)
        tk.Label(welcome, text="Módulo de Reportes",
                 font=('Segoe UI', 24, 'bold'), bg='white', fg='#2c3e50').pack()
        tk.Label(welcome, text="Selecciona un tipo de reporte del menú lateral",
                 font=('Segoe UI', 12), bg='white', fg='#7f8c8d').pack(pady=10)

    # ==========================================
    # UTILIDADES COMUNES
    # ==========================================

    def _clear_right_panel(self):
        for widget in self.right_panel.winfo_children():
            widget.destroy()
        # Resetear scroll al inicio
        self._right_canvas.yview_moveto(0)

    def _create_report_header(self, title, icon, color):
        header = tk.Frame(self.right_panel, bg=color, height=60)
        header.pack(fill='x')
        header.pack_propagate(False)
        tk.Label(header, text=f"{icon} {title}",
                 font=('Segoe UI', 16, 'bold'), bg=color, fg='white'
                 ).pack(side='left', padx=20, pady=15)
        return header

    def create_summary_card(self, parent, title, value, icon, color):
        """Crear tarjeta de resumen"""
        card = tk.Frame(parent, bg=color, relief='raised', borderwidth=2)
        card.pack(side='left', fill='both', expand=True, padx=5)
        
        tk.Label(card, text=icon, font=('Segoe UI', 24), bg=color, fg='white').pack(pady=5)
        tk.Label(card, text=str(value), font=('Segoe UI', 16, 'bold'), bg=color, fg='white').pack()
        tk.Label(card, text=title, font=('Segoe UI', 9), bg=color, fg='white').pack(pady=5)

    def _create_treeview(self, parent, columns, height=12):
        """Crear Treeview con scrollbar"""
        tree_frame = tk.Frame(parent, bg='white')
        tree_frame.pack(fill='both', expand=True, pady=5)
        
        scroll_y = tk.Scrollbar(tree_frame)
        scroll_y.pack(side='right', fill='y')
        
        scroll_x = tk.Scrollbar(tree_frame, orient='horizontal')
        scroll_x.pack(side='bottom', fill='x')
        
        tree = ttk.Treeview(
            tree_frame, columns=columns, show='headings',
            height=height, yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set
        )
        
        for col in columns:
            tree.heading(col, text=col)
            tree.column(col, width=110, minwidth=80)
        
        tree.pack(fill='both', expand=True)
        scroll_y.config(command=tree.yview)
        scroll_x.config(command=tree.xview)
        
        return tree

    def _create_export_button(self, parent, command, text="📥 Exportar a Excel"):
        tk.Button(
            parent, text=text, command=command,
            font=('Segoe UI', 10, 'bold'), bg='#27ae60', fg='white',
            relief='flat', padx=15, pady=8, cursor='hand2'
        ).pack(pady=10)

    def _embed_chart(self, parent, figsize=(6, 3)):
        """Crear un gráfico matplotlib embebido en un frame Tkinter.
        Retorna (fig, ax, canvas) para que el llamador dibuje su gráfico."""
        chart_frame = tk.Frame(parent, bg='white', relief='groove', borderwidth=1)
        chart_frame.pack(fill='x', pady=8, padx=5)

        fig = Figure(figsize=figsize, dpi=90, facecolor='white')
        ax = fig.add_subplot(111)
        ax.set_facecolor('#fafafa')

        canvas = FigureCanvasTkAgg(fig, master=chart_frame)
        canvas.get_tk_widget().pack(fill='x', expand=True)

        return fig, ax, canvas

    def _embed_chart_dual(self, parent, figsize=(7, 3)):
        """Crear un gráfico con dos subplots lado a lado."""
        chart_frame = tk.Frame(parent, bg='white', relief='groove', borderwidth=1)
        chart_frame.pack(fill='x', pady=8, padx=5)

        fig = Figure(figsize=figsize, dpi=90, facecolor='white')
        ax1 = fig.add_subplot(121)
        ax2 = fig.add_subplot(122)
        ax1.set_facecolor('#fafafa')
        ax2.set_facecolor('#fafafa')

        canvas = FigureCanvasTkAgg(fig, master=chart_frame)
        canvas.get_tk_widget().pack(fill='x', expand=True)

        return fig, ax1, ax2, canvas

    def _export_to_excel(self, data_rows, columns, sheet_name, title):
        """Exportar datos a Excel usando openpyxl"""
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            messagebox.showerror("Error", "Se requiere openpyxl. Instale con: pip install openpyxl")
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"Reporte_{sheet_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not filepath:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name[:31]

            # Titulo
            header_fill = PatternFill(start_color='8E44AD', end_color='8E44AD', fill_type='solid')
            header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
            ws.cell(row=1, column=1, value=title).font = Font(name='Segoe UI', size=14, bold=True)
            ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')
            
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
            ws.cell(row=2, column=1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = Font(size=9, italic=True)
            ws.cell(row=2, column=1).alignment = Alignment(horizontal='center')

            # Encabezados
            for col_idx, col_name in enumerate(columns, 1):
                cell = ws.cell(row=4, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')

            # Datos
            for row_idx, row_data in enumerate(data_rows, 5):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.alignment = Alignment(horizontal='center')

            # Auto-ancho
            for col_idx in range(1, len(columns) + 1):
                max_len = max(
                    len(str(ws.cell(row=r, column=col_idx).value or ''))
                    for r in range(4, ws.max_row + 1)
                )
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 30)

            wb.save(filepath)
            messagebox.showinfo("Exportado", f"Reporte guardado en:\n{filepath}")
            os.startfile(os.path.dirname(filepath))
        except Exception as e:
            messagebox.showerror("Error", f"Error al exportar: {str(e)}")

    # ==========================================
    # 1. REPORTE DE VENTAS (existente)
    # ==========================================

    def show_sales_report(self):
        """Mostrar reporte de ventas"""
        self._clear_right_panel()
        self._create_report_header("REPORTE DE VENTAS", "📈", "#3498db")
        
        filters = tk.Frame(self.right_panel, bg='white')
        filters.pack(fill='x', padx=20, pady=15)
        
        date_frame = tk.Frame(filters, bg='white')
        date_frame.pack(fill='x', pady=5)
        
        tk.Label(date_frame, text="Desde:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        self.sales_start_date = DateEntry(date_frame, width=12, background='#3498db',
                                           foreground='white', borderwidth=2, date_pattern='dd/mm/yyyy')
        self.sales_start_date.pack(side='left', padx=5)
        
        tk.Label(date_frame, text="Hasta:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        self.sales_end_date = DateEntry(date_frame, width=12, background='#3498db',
                                         foreground='white', borderwidth=2, date_pattern='dd/mm/yyyy')
        self.sales_end_date.pack(side='left', padx=5)
        
        tk.Button(
            date_frame, text="🔍 Generar Reporte", command=self.generate_sales_report,
            font=('Segoe UI', 10, 'bold'), bg='#27ae60', fg='white',
            relief='flat', padx=15, pady=5, cursor='hand2'
        ).pack(side='left', padx=10)
        
        self.sales_results = tk.Frame(self.right_panel, bg='white')
        self.sales_results.pack(fill='both', expand=True, padx=20, pady=10)
    
    def generate_sales_report(self):
        """Generar reporte de ventas"""
        try:
            start = self.sales_start_date.get_date().strftime('%Y-%m-%d')
            end = self.sales_end_date.get_date().strftime('%Y-%m-%d')
            
            result = self.controller.get_sales_report(start, end)
            
            if not result['success']:
                messagebox.showerror("Error", result['message'])
                return
            
            for widget in self.sales_results.winfo_children():
                widget.destroy()
            
            data = result['data']
            summary = data['summary']
            
            cards_frame = tk.Frame(self.sales_results, bg='white')
            cards_frame.pack(fill='x', pady=10)
            
            cards = [
                ('Total Ventas', summary['total_sales'], '🛒', '#3498db'),
                ('Monto Total', f"S/ {summary['total_amount']:.2f}", '💰', '#27ae60'),
                ('Ticket Promedio', f"S/ {summary['average_ticket']:.2f}", '🎫', '#f39c12'),
                ('Descuentos', f"S/ {summary['total_discount']:.2f}", '🏷️', '#e74c3c')
            ]
            
            for title, value, icon, color in cards:
                self.create_summary_card(cards_frame, title, value, icon, color)
            
            # === GRÁFICO: Distribución por Método de Pago ===
            try:
                pay_counts = defaultdict(float)
                for sale in data['sales']:
                    method = sale['payment_method'].upper()
                    pay_counts[method] += float(sale['total_amount'])
                if pay_counts:
                    fig, ax, canvas = self._embed_chart(self.sales_results, figsize=(5, 2.5))
                    labels = list(pay_counts.keys())
                    sizes = list(pay_counts.values())
                    colors_pie = ['#3498db', '#27ae60', '#f39c12', '#e74c3c', '#9b59b6', '#1abc9c']
                    wedges, texts, autotexts = ax.pie(
                        sizes, labels=labels, autopct='%1.1f%%',
                        colors=colors_pie[:len(labels)], startangle=90,
                        textprops={'fontsize': 8}
                    )
                    ax.set_title('Distribución por Método de Pago', fontsize=10, fontweight='bold', pad=10)
                    fig.tight_layout()
                    canvas.draw()
            except Exception:
                pass
            
            columns = ('Número', 'Fecha', 'Cajero', 'Cliente', 'Total', 'Método Pago')
            tree = self._create_treeview(self.sales_results, columns)
            
            for sale in data['sales']:
                sale_date = sale['sale_date']
                if isinstance(sale_date, datetime):
                    sale_date = sale_date.strftime('%d/%m/%Y %H:%M')
                tree.insert('', 'end', values=(
                    sale['sale_number'],
                    sale_date,
                    sale['cashier_name'] or 'N/A',
                    sale['customer_name'] or 'Cliente Genérico',
                    f"S/ {sale['total_amount']:.2f}",
                    sale['payment_method'].upper()
                ))
            
            # Exportar
            self._sales_data = data
            self._create_export_button(self.sales_results, self._export_sales_report)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error al generar reporte: {str(e)}")

    def _export_sales_report(self):
        data = self._sales_data
        columns = ['Número', 'Fecha', 'Cajero', 'Cliente', 'Total', 'Método Pago']
        rows = []
        for s in data['sales']:
            sd = s['sale_date']
            if isinstance(sd, datetime):
                sd = sd.strftime('%d/%m/%Y %H:%M')
            rows.append([
                s['sale_number'], sd, s['cashier_name'] or 'N/A',
                s['customer_name'] or 'General',
                round(float(s['total_amount']), 2),
                s['payment_method'].upper()
            ])
        period = data.get('period', {})
        title = f"Reporte de Ventas ({period.get('start', '')} - {period.get('end', '')})"
        self._export_to_excel(rows, columns, 'Ventas', title)

    # ==========================================
    # 2. FLUJO DE PRODUCTOS
    # ==========================================

    def show_products_report(self):
        """Mostrar reporte de flujo de productos"""
        self._clear_right_panel()
        self._create_report_header("FLUJO DE PRODUCTOS", "📦", "#e67e22")
        
        filters = tk.Frame(self.right_panel, bg='white')
        filters.pack(fill='x', padx=20, pady=15)
        
        date_frame = tk.Frame(filters, bg='white')
        date_frame.pack(fill='x', pady=5)
        
        tk.Label(date_frame, text="Desde:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        self.flow_start_date = DateEntry(date_frame, width=12, background='#e67e22',
                                          foreground='white', borderwidth=2, date_pattern='dd/mm/yyyy')
        self.flow_start_date.set_date(datetime.now() - timedelta(days=30))
        self.flow_start_date.pack(side='left', padx=5)
        
        tk.Label(date_frame, text="Hasta:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        self.flow_end_date = DateEntry(date_frame, width=12, background='#e67e22',
                                        foreground='white', borderwidth=2, date_pattern='dd/mm/yyyy')
        self.flow_end_date.pack(side='left', padx=5)
        
        tk.Button(
            date_frame, text="🔍 Generar Reporte", command=self._generate_flow_report,
            font=('Segoe UI', 10, 'bold'), bg='#27ae60', fg='white',
            relief='flat', padx=15, pady=5, cursor='hand2'
        ).pack(side='left', padx=10)
        
        self.flow_results = tk.Frame(self.right_panel, bg='white')
        self.flow_results.pack(fill='both', expand=True, padx=20, pady=10)

    def _generate_flow_report(self):
        try:
            start = self.flow_start_date.get_date().strftime('%Y-%m-%d')
            end = self.flow_end_date.get_date().strftime('%Y-%m-%d')
            
            result = self.controller.get_product_flow_report(start, end)
            
            if not result['success']:
                messagebox.showerror("Error", result['message'])
                return
            
            for widget in self.flow_results.winfo_children():
                widget.destroy()
            
            data = result['data']
            summary = data['summary']
            
            cards_frame = tk.Frame(self.flow_results, bg='white')
            cards_frame.pack(fill='x', pady=10)
            
            cards = [
                ('Productos Activos', summary['total_products'], '📦', '#3498db'),
                ('Con Movimiento', summary['products_with_movement'], '🔄', '#e67e22'),
                ('Total Entradas', f"+{summary['total_entradas']}", '📥', '#27ae60'),
                ('Total Salidas', f"-{summary['total_salidas']}", '📤', '#e74c3c'),
            ]
            
            for title, value, icon, color in cards:
                self.create_summary_card(cards_frame, title, value, icon, color)
            
            # === GRÁFICO: Top 10 Productos con Movimiento ===
            try:
                prods_mov = [p for p in data['products'] if int(p['total_movimientos'] or 0) > 0]
                prods_mov.sort(key=lambda x: int(x['total_movimientos'] or 0), reverse=True)
                top10 = prods_mov[:10]
                if top10:
                    fig, ax, canvas = self._embed_chart(self.flow_results, figsize=(6, 3))
                    names = [p['name'][:18] for p in reversed(top10)]
                    entradas = [int(p['total_entradas']) for p in reversed(top10)]
                    salidas = [int(p['total_salidas']) for p in reversed(top10)]
                    y_pos = range(len(names))
                    ax.barh(y_pos, entradas, height=0.4, label='Entradas', color='#27ae60', align='center')
                    ax.barh([y + 0.4 for y in y_pos], salidas, height=0.4, label='Salidas', color='#e74c3c', align='center')
                    ax.set_yticks([y + 0.2 for y in y_pos])
                    ax.set_yticklabels(names, fontsize=7)
                    ax.set_xlabel('Unidades', fontsize=8)
                    ax.set_title('Top 10 Productos con Mayor Movimiento', fontsize=10, fontweight='bold')
                    ax.legend(fontsize=7, loc='lower right')
                    ax.tick_params(axis='x', labelsize=7)
                    fig.tight_layout()
                    canvas.draw()
            except Exception:
                pass
            
            columns = ('SKU', 'Producto', 'Categoría', 'Entradas', 'Salidas', 'Stock Actual')
            tree = self._create_treeview(self.flow_results, columns)
            
            for p in data['products']:
                if int(p['total_movimientos'] or 0) > 0:
                    tree.insert('', 'end', values=(
                        p['sku'], p['name'], p['category_name'] or 'S/C',
                        f"+{int(p['total_entradas'])}", f"-{int(p['total_salidas'])}",
                        int(p['stock_quantity'] or 0)
                    ))
            
            self._flow_data = data
            self._create_export_button(self.flow_results, self._export_flow_report)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error: {str(e)}")

    def _export_flow_report(self):
        data = self._flow_data
        columns = ['SKU', 'Producto', 'Categoría', 'Entradas', 'Salidas', 'Stock Actual']
        rows = []
        for p in data['products']:
            if int(p['total_movimientos'] or 0) > 0:
                rows.append([
                    p['sku'], p['name'], p['category_name'] or 'S/C',
                    int(p['total_entradas']), int(p['total_salidas']),
                    int(p['stock_quantity'] or 0)
                ])
        period = data.get('period', {})
        self._export_to_excel(rows, columns, 'Flujo_Productos',
                              f"Flujo de Productos ({period.get('start', '')} - {period.get('end', '')})")

    # ==========================================
    # 3. VENTAS MENSUALES
    # ==========================================

    def show_monthly_sales_report(self):
        """Mostrar reporte mensual"""
        self._clear_right_panel()
        self._create_report_header("VENTAS MENSUALES", "📅", "#2980b9")
        
        filters = tk.Frame(self.right_panel, bg='white')
        filters.pack(fill='x', padx=20, pady=15)
        
        year_frame = tk.Frame(filters, bg='white')
        year_frame.pack(fill='x', pady=5)
        
        tk.Label(year_frame, text="Año:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        
        current_year = datetime.now().year
        self.year_var = tk.StringVar(value=str(current_year))
        year_combo = ttk.Combobox(year_frame, textvariable=self.year_var, width=8,
                                   values=[str(y) for y in range(current_year - 5, current_year + 1)],
                                   state='readonly')
        year_combo.pack(side='left', padx=5)
        
        tk.Button(
            year_frame, text="🔍 Generar Reporte", command=self._generate_monthly_report,
            font=('Segoe UI', 10, 'bold'), bg='#27ae60', fg='white',
            relief='flat', padx=15, pady=5, cursor='hand2'
        ).pack(side='left', padx=10)
        
        self.monthly_results = tk.Frame(self.right_panel, bg='white')
        self.monthly_results.pack(fill='both', expand=True, padx=20, pady=10)

    def _generate_monthly_report(self):
        try:
            year = int(self.year_var.get())
            result = self.controller.get_monthly_sales_report(year)
            
            if not result['success']:
                messagebox.showerror("Error", result['message'])
                return
            
            for widget in self.monthly_results.winfo_children():
                widget.destroy()
            
            data = result['data']
            summary = data['summary']
            
            cards_frame = tk.Frame(self.monthly_results, bg='white')
            cards_frame.pack(fill='x', pady=10)
            
            cards = [
                ('Ventas Totales', summary['total_ventas'], '🛒', '#3498db'),
                ('Monto Anual', f"S/ {summary['monto_total']:.2f}", '💰', '#27ae60'),
                ('Ticket Promedio', f"S/ {summary['ticket_promedio']:.2f}", '🎫', '#f39c12'),
                ('Mejor Mes', summary['mejor_mes'], '⭐', '#9b59b6'),
            ]
            
            for title, value, icon, color in cards:
                self.create_summary_card(cards_frame, title, value, icon, color)
            
            # === GRÁFICO: Ventas Mensuales (Barras + Línea) ===
            try:
                meses = [m['nombre'][:3] for m in data['months']]
                montos = [m['monto_total'] for m in data['months']]
                tickets = [m['ticket_promedio'] for m in data['months']]
                if any(m > 0 for m in montos):
                    fig, ax1, canvas = self._embed_chart(self.monthly_results, figsize=(7, 3))
                    x = range(len(meses))
                    bars = ax1.bar(x, montos, color='#3498db', alpha=0.8, label='Monto Total (S/)')
                    ax1.set_xlabel('Mes', fontsize=8)
                    ax1.set_ylabel('Monto (S/)', fontsize=8, color='#3498db')
                    ax1.set_xticks(list(x))
                    ax1.set_xticklabels(meses, fontsize=7)
                    ax1.tick_params(axis='y', labelsize=7, labelcolor='#3498db')
                    # Línea de ticket promedio en segundo eje
                    ax2 = ax1.twinx()
                    ax2.plot(list(x), tickets, color='#e74c3c', marker='o', linewidth=2,
                             markersize=4, label='Ticket Prom. (S/)')
                    ax2.set_ylabel('Ticket Promedio (S/)', fontsize=8, color='#e74c3c')
                    ax2.tick_params(axis='y', labelsize=7, labelcolor='#e74c3c')
                    ax1.set_title(f'Ventas Mensuales - {data["year"]}', fontsize=10, fontweight='bold')
                    # Leyenda combinada
                    lines1, labels1 = ax1.get_legend_handles_labels()
                    lines2, labels2 = ax2.get_legend_handles_labels()
                    ax1.legend(lines1 + lines2, labels1 + labels2, fontsize=7, loc='upper left')
                    fig.tight_layout()
                    canvas.draw()
            except Exception:
                pass
            
            columns = ('Mes', 'Ventas', 'Monto Total', 'Ticket Prom.', 'Productos', 'Descuentos')
            tree = self._create_treeview(self.monthly_results, columns)
            
            for m in data['months']:
                tree.insert('', 'end', values=(
                    m['nombre'],
                    m['total_ventas'],
                    f"S/ {m['monto_total']:.2f}",
                    f"S/ {m['ticket_promedio']:.2f}",
                    m['total_productos'],
                    f"S/ {m['total_descuentos']:.2f}"
                ))
            
            self._monthly_data = data
            self._create_export_button(self.monthly_results, self._export_monthly_report)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error: {str(e)}")

    def _export_monthly_report(self):
        data = self._monthly_data
        columns = ['Mes', 'Ventas', 'Monto Total', 'Ticket Promedio', 'Productos', 'Descuentos']
        rows = []
        for m in data['months']:
            rows.append([
                m['nombre'], m['total_ventas'],
                round(m['monto_total'], 2), round(m['ticket_promedio'], 2),
                m['total_productos'], round(m['total_descuentos'], 2)
            ])
        self._export_to_excel(rows, columns, 'Ventas_Mensuales',
                              f"Ventas Mensuales - Año {data['year']}")

    # ==========================================
    # 4. INVERSIÓN EN INVENTARIO
    # ==========================================

    def show_inventory_report(self):
        """Mostrar reporte de inversión en inventario"""
        self._clear_right_panel()
        self._create_report_header("INVERSIÓN EN INVENTARIO", "💰", "#16a085")
        
        # Info
        info_frame = tk.Frame(self.right_panel, bg='#eaf6f4')
        info_frame.pack(fill='x', padx=20, pady=10)
        tk.Label(
            info_frame,
            text="💡 Este reporte muestra cuánto dinero tienes invertido en productos dentro de tu tienda",
            font=('Segoe UI', 10), bg='#eaf6f4', fg='#16a085', wraplength=600
        ).pack(padx=15, pady=10)
        
        # Botón generar
        btn_frame = tk.Frame(self.right_panel, bg='white')
        btn_frame.pack(fill='x', padx=20)
        tk.Button(
            btn_frame, text="🔍 Calcular Inversión", command=self._generate_inventory_report,
            font=('Segoe UI', 10, 'bold'), bg='#16a085', fg='white',
            relief='flat', padx=15, pady=8, cursor='hand2'
        ).pack(pady=5)
        
        self.inventory_results = tk.Frame(self.right_panel, bg='white')
        self.inventory_results.pack(fill='both', expand=True, padx=20, pady=10)

    def _generate_inventory_report(self):
        try:
            result = self.controller.get_inventory_investment_report()
            
            if not result['success']:
                messagebox.showerror("Error", result['message'])
                return
            
            for widget in self.inventory_results.winfo_children():
                widget.destroy()
            
            data = result['data']
            summary = data['summary']
            
            # Cards principales
            cards_frame = tk.Frame(self.inventory_results, bg='white')
            cards_frame.pack(fill='x', pady=10)
            
            cards = [
                ('Inversión (Costo)', f"S/ {summary['total_inversion_costo']:.2f}", '💵', '#e74c3c'),
                ('Valor de Venta', f"S/ {summary['total_valor_venta']:.2f}", '💰', '#27ae60'),
                ('Ganancia Potencial', f"S/ {summary['ganancia_potencial']:.2f}", '📈', '#f39c12'),
                ('Margen Global', f"{summary['margen_global']:.1f}%", '📊', '#3498db'),
            ]
            
            for title, value, icon, color in cards:
                self.create_summary_card(cards_frame, title, value, icon, color)
            
            # Info extra
            info2 = tk.Frame(self.inventory_results, bg='white')
            info2.pack(fill='x', pady=5)
            tk.Label(
                info2,
                text=f"📦 {summary['total_products']} productos  |  "
                     f"🔢 {summary['total_unidades']} unidades en stock",
                font=('Segoe UI', 11), bg='white', fg='#2c3e50'
            ).pack(pady=5)
            
            # === GRÁFICO: Inversión por Categoría (Pie) ===
            try:
                cats_with_data = [c for c in data['categories'] if float(c['inversion_costo'] or 0) > 0]
                if cats_with_data:
                    fig, ax1, ax2, canvas = self._embed_chart_dual(self.inventory_results, figsize=(7, 3))
                    cat_names = [c['category_name'][:15] for c in cats_with_data]
                    cat_inv = [float(c['inversion_costo'] or 0) for c in cats_with_data]
                    cat_venta = [float(c['valor_venta'] or 0) for c in cats_with_data]
                    colors_pie = ['#3498db', '#27ae60', '#f39c12', '#e74c3c', '#9b59b6',
                                  '#1abc9c', '#e67e22', '#2c3e50', '#d35400', '#c0392b']
                    ax1.pie(cat_inv, labels=cat_names, autopct='%1.1f%%',
                            colors=colors_pie[:len(cat_names)], startangle=90,
                            textprops={'fontsize': 7})
                    ax1.set_title('Inversión (Costo)', fontsize=9, fontweight='bold')
                    ax2.pie(cat_venta, labels=cat_names, autopct='%1.1f%%',
                            colors=colors_pie[:len(cat_names)], startangle=90,
                            textprops={'fontsize': 7})
                    ax2.set_title('Valor de Venta', fontsize=9, fontweight='bold')
                    fig.tight_layout()
                    canvas.draw()
            except Exception:
                pass
            
            # Tabla por categoría
            tk.Label(self.inventory_results, text="Inversión por Categoría",
                     font=('Segoe UI', 12, 'bold'), bg='white', fg='#2c3e50'
                     ).pack(anchor='w', pady=(10, 0))
            
            cat_cols = ('Categoría', 'Productos', 'Unidades', 'Inversión (Costo)', 'Valor Venta')
            cat_tree = self._create_treeview(self.inventory_results, cat_cols, height=6)
            
            for c in data['categories']:
                cat_tree.insert('', 'end', values=(
                    c['category_name'],
                    int(c['total_productos']),
                    int(c['total_unidades'] or 0),
                    f"S/ {float(c['inversion_costo'] or 0):.2f}",
                    f"S/ {float(c['valor_venta'] or 0):.2f}"
                ))
            
            # Tabla por producto (Top 20)
            tk.Label(self.inventory_results, text="Top 20 Productos por Inversión",
                     font=('Segoe UI', 12, 'bold'), bg='white', fg='#2c3e50'
                     ).pack(anchor='w', pady=(10, 0))
            
            prod_cols = ('SKU', 'Producto', 'Stock', 'Costo Unit.', 'Inversión', 'Valor Venta', 'Margen')
            prod_tree = self._create_treeview(self.inventory_results, prod_cols, height=8)
            
            for p in data['products'][:20]:
                prod_tree.insert('', 'end', values=(
                    p['sku'], p['name'],
                    int(p['stock_quantity'] or 0),
                    f"S/ {float(p['cost'] or 0):.2f}",
                    f"S/ {float(p['inversion_costo'] or 0):.2f}",
                    f"S/ {float(p['valor_venta'] or 0):.2f}",
                    f"{float(p['margen_pct'] or 0):.1f}%"
                ))
            
            self._inventory_data = data
            self._create_export_button(self.inventory_results, self._export_inventory_report)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error: {str(e)}")

    def _export_inventory_report(self):
        data = self._inventory_data
        columns = ['SKU', 'Producto', 'Categoría', 'Stock', 'Costo Unit.', 'Precio Venta',
                    'Inversión (Costo)', 'Valor Venta', 'Margen %']
        rows = []
        for p in data['products']:
            rows.append([
                p['sku'], p['name'], p['category_name'] or 'S/C',
                int(p['stock_quantity'] or 0),
                round(float(p['cost'] or 0), 2),
                round(float(p['price'] or 0), 2),
                round(float(p['inversion_costo'] or 0), 2),
                round(float(p['valor_venta'] or 0), 2),
                round(float(p['margen_pct'] or 0), 2)
            ])
        summary = data['summary']
        self._export_to_excel(rows, columns, 'Inversion_Inventario',
                              f"Inversión en Inventario - Total: S/ {summary['total_inversion_costo']:.2f}")

    # ==========================================
    # 5. PRODUCTO ESPECÍFICO
    # ==========================================

    def show_specific_product_report(self):
        """Mostrar reporte de producto específico"""
        self._clear_right_panel()
        self._create_report_header("PRODUCTO ESPECÍFICO", "🔍", "#8e44ad")
        
        filters = tk.Frame(self.right_panel, bg='white')
        filters.pack(fill='x', padx=20, pady=15)
        
        # Selector de producto
        prod_frame = tk.Frame(filters, bg='white')
        prod_frame.pack(fill='x', pady=5)
        
        tk.Label(prod_frame, text="Producto:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        
        self.product_list = self.controller.get_all_products_list()
        product_names = [f"{p['sku']} - {p['name']}" for p in self.product_list]
        
        self.product_var = tk.StringVar()
        self.product_combo = ttk.Combobox(
            prod_frame, textvariable=self.product_var,
            values=product_names, width=40, state='readonly'
        )
        self.product_combo.pack(side='left', padx=5)
        if product_names:
            self.product_combo.current(0)
        
        # Filtros de fecha
        date_frame = tk.Frame(filters, bg='white')
        date_frame.pack(fill='x', pady=5)
        
        tk.Label(date_frame, text="Desde:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        self.prod_start_date = DateEntry(date_frame, width=12, background='#8e44ad',
                                          foreground='white', borderwidth=2, date_pattern='dd/mm/yyyy')
        self.prod_start_date.set_date(datetime.now() - timedelta(days=30))
        self.prod_start_date.pack(side='left', padx=5)
        
        tk.Label(date_frame, text="Hasta:", font=('Segoe UI', 10), bg='white').pack(side='left', padx=5)
        self.prod_end_date = DateEntry(date_frame, width=12, background='#8e44ad',
                                        foreground='white', borderwidth=2, date_pattern='dd/mm/yyyy')
        self.prod_end_date.pack(side='left', padx=5)
        
        tk.Button(
            date_frame, text="🔍 Analizar Producto", command=self._generate_specific_report,
            font=('Segoe UI', 10, 'bold'), bg='#27ae60', fg='white',
            relief='flat', padx=15, pady=5, cursor='hand2'
        ).pack(side='left', padx=10)
        
        self.specific_results = tk.Frame(self.right_panel, bg='white')
        self.specific_results.pack(fill='both', expand=True, padx=20, pady=10)

    def _generate_specific_report(self):
        try:
            idx = self.product_combo.current()
            if idx < 0:
                messagebox.showwarning("Aviso", "Seleccione un producto")
                return
            
            product_id = self.product_list[idx]['id']
            start = self.prod_start_date.get_date().strftime('%Y-%m-%d')
            end = self.prod_end_date.get_date().strftime('%Y-%m-%d')
            
            result = self.controller.get_specific_product_report(product_id, start, end)
            
            if not result['success']:
                messagebox.showerror("Error", result['message'])
                return
            
            for widget in self.specific_results.winfo_children():
                widget.destroy()
            
            data = result['data']
            summary = data['summary']
            product = data['product']
            
            # Info del producto
            info_frame = tk.Frame(self.specific_results, bg='#f8f9fa', relief='solid', borderwidth=1)
            info_frame.pack(fill='x', pady=5)
            tk.Label(
                info_frame,
                text=f"📦 {product['name']}  |  SKU: {product['sku']}  |  "
                     f"Categoría: {product.get('category_name', 'S/C')}  |  "
                     f"Costo: S/ {float(product['cost'] or 0):.2f}  |  "
                     f"Precio: S/ {float(product['price'] or 0):.2f}",
                font=('Segoe UI', 10), bg='#f8f9fa', fg='#2c3e50', wraplength=700
            ).pack(padx=10, pady=8)
            
            # Cards de KPIs
            cards_frame = tk.Frame(self.specific_results, bg='white')
            cards_frame.pack(fill='x', pady=10)
            
            cards = [
                ('Uds. Vendidas', summary['total_unidades_vendidas'], '📦', '#3498db'),
                ('Ingresos', f"S/ {summary['total_ingresos']:.2f}", '💰', '#27ae60'),
                ('Ganancia', f"S/ {summary['ganancia']:.2f}", '📈', '#f39c12'),
                ('Stock Actual', summary['stock_actual'], '🔢', '#9b59b6'),
            ]
            
            for title, value, icon, color in cards:
                self.create_summary_card(cards_frame, title, value, icon, color)
            
            # Info adicional
            extra = tk.Frame(self.specific_results, bg='white')
            extra.pack(fill='x', pady=5)
            tk.Label(
                extra,
                text=f"🛒 Vendido {summary['veces_vendido']} veces  |  "
                     f"Precio Prom: S/ {summary['precio_promedio']:.2f}  |  "
                     f"Costo Total: S/ {summary['costo_total_vendido']:.2f}",
                font=('Segoe UI', 10), bg='white', fg='#7f8c8d'
            ).pack(pady=3)
            
            # === GRÁFICO: Unidades vendidas por fecha ===
            try:
                ventas_por_fecha = defaultdict(int)
                for s in data['sales']:
                    sd = s['sale_date']
                    if isinstance(sd, datetime):
                        key = sd.strftime('%d/%m')
                    else:
                        key = str(sd)[:10]
                    ventas_por_fecha[key] += int(s['quantity'])
                if ventas_por_fecha:
                    fig, ax, canvas = self._embed_chart(self.specific_results, figsize=(6, 2.5))
                    fechas = list(ventas_por_fecha.keys())
                    cantidades = list(ventas_por_fecha.values())
                    ax.bar(range(len(fechas)), cantidades, color='#8e44ad', alpha=0.85)
                    ax.set_xticks(range(len(fechas)))
                    ax.set_xticklabels(fechas, fontsize=7, rotation=45, ha='right')
                    ax.set_ylabel('Unidades', fontsize=8)
                    ax.set_title(f'Unidades Vendidas por Fecha - {product["name"]}',
                                 fontsize=9, fontweight='bold')
                    ax.tick_params(axis='y', labelsize=7)
                    fig.tight_layout()
                    canvas.draw()
            except Exception:
                pass
            
            # Tabla de ventas
            tk.Label(self.specific_results, text="Historial de Ventas",
                     font=('Segoe UI', 12, 'bold'), bg='white', fg='#2c3e50'
                     ).pack(anchor='w', pady=(10, 0))
            
            sale_cols = ('Venta', 'Fecha', 'Cantidad', 'Precio Unit.', 'Subtotal', 'Cajero')
            sale_tree = self._create_treeview(self.specific_results, sale_cols, height=8)
            
            for s in data['sales']:
                sd = s['sale_date']
                if isinstance(sd, datetime):
                    sd = sd.strftime('%d/%m/%Y %H:%M')
                sale_tree.insert('', 'end', values=(
                    s['sale_number'], sd,
                    int(s['quantity']),
                    f"S/ {float(s['unit_price'] or 0):.2f}",
                    f"S/ {float(s['subtotal'] or 0):.2f}",
                    s['cashier_name'] or 'N/A'
                ))
            
            self._specific_data = data
            self._create_export_button(self.specific_results, self._export_specific_report)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error: {str(e)}")

    def _export_specific_report(self):
        data = self._specific_data
        product = data['product']
        columns = ['Venta', 'Fecha', 'Cantidad', 'Precio Unit.', 'Subtotal', 'Cajero', 'Cliente']
        rows = []
        for s in data['sales']:
            sd = s['sale_date']
            if isinstance(sd, datetime):
                sd = sd.strftime('%d/%m/%Y %H:%M')
            rows.append([
                s['sale_number'], sd, int(s['quantity']),
                round(float(s['unit_price'] or 0), 2),
                round(float(s['subtotal'] or 0), 2),
                s['cashier_name'] or 'N/A',
                s['customer_name'] or 'General'
            ])
        period = data.get('period', {})
        self._export_to_excel(
            rows, columns, f"Producto_{product['sku']}",
            f"Reporte - {product['name']} ({period.get('start', '')} - {period.get('end', '')})"
        )

    # ==========================================
    # MÉTODOS DE CICLO DE VIDA
    # ==========================================

    def show(self):
        """Mostrar vista"""
        self.main_frame.pack(fill='both', expand=True)
    
    def hide(self):
        """Ocultar vista"""
        self.main_frame.pack_forget()
    
    def destroy(self):
        """Destruir vista"""
        self.main_frame.destroy()
